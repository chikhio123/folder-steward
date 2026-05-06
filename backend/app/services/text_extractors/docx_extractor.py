from pathlib import Path
import docx
from docx.opc.exceptions import PackageNotFoundError
from .base import TextExtractor, ExtractResult

class DocxExtractor(TextExtractor):
    supported_extensions = {".docx"}

    def extract(self, path: Path) -> ExtractResult:
        # 限制最大读取大小，例如 50MB
        MAX_SIZE = 50 * 1024 * 1024
        MAX_CHARS = 1_000_000

        stat = path.stat()
        if stat.st_size > MAX_SIZE:
            return ExtractResult(
                text="",
                warnings=[f"File too large ({stat.st_size} bytes), skipped."]
            )

        text_blocks = []
        warnings = []
        metadata = {}

        try:
            doc = docx.Document(str(path))
            current_length = 0
            truncated = False

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    text_blocks.append(text)
                    current_length += len(text)
                    if current_length > MAX_CHARS:
                        truncated = True
                        break

            # If not yet truncated, extract tables
            if not truncated:
                for table in doc.tables:
                    if truncated:
                        break
                    for row in table.rows:
                        row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_data:
                            joined_row = " | ".join(row_data)
                            text_blocks.append(joined_row)
                            current_length += len(joined_row)
                            if current_length > MAX_CHARS:
                                truncated = True
                                break

            full_text = "\n".join(text_blocks).strip()

            if not full_text:
                return ExtractResult(
                    text="",
                    warnings=["no_extractable_text"]
                )

            if len(full_text) > MAX_CHARS:
                full_text = full_text[:MAX_CHARS]

            if truncated:
                warnings.append(f"Text truncated to {MAX_CHARS} characters during parsing.")

            return ExtractResult(text=full_text, metadata=metadata, warnings=warnings)

        except PackageNotFoundError:
            return ExtractResult(text="", warnings=["File is not a valid DOCX package."])
        except Exception as e:
            return ExtractResult(text="", warnings=[f"Unexpected error reading DOCX: {e}"])
